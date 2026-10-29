"""Convert the Danish Food Composition Database (Frida / FCDB) into the compact
file the integration ships with: custom_components/hass_recipes/data/frida.json.

Download the dataset from DTU Data (https://doi.org/10.11583/DTU.32312844), unzip it,
and run:

    pip install openpyxl
    python dev/frida_convert.py path/to/FCDB_6.1_Dataset.xlsx 6.1 "May 2026"

The dataset is published under CC BY 4.0; the attribution DTU asks for is stored in the
output file and shown in the panel wherever nutrition is displayed.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).parents[1] / "custom_components/hass_recipes/data/frida.json"

# output key -> Frida ParameterID (all per 100 g edible part)
PARAMS = {
    "kcal": 356,      # Energy (kcal)
    "protein": 218,   # Protein
    "fat": 141,       # Fat
    "carbs": 172,     # Available carbohydrates
    "fibre": 168,     # Dietary fibre
    "sugars": 245,    # Sum sugars
    "sat_fat": 248,   # Sum saturated fatty acids
    "salt": 327,      # Salt labelling
}
FIELDS = ["id", "name_da", "name_en", "group", *PARAMS]


def _num(value) -> float | None:
    if value in (None, "", "NULL"):
        return None
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def convert(xlsx: Path, version: str, date: str) -> dict:
    wb = openpyxl.load_workbook(xlsx, read_only=True)

    groups: dict[int, str] = {}
    for row in wb["Food"].iter_rows(min_row=2, values_only=True):
        if row[2] is not None:
            groups[int(row[2])] = row[11] or ""

    rows = wb["Data_Table"].iter_rows(values_only=True)
    header = None
    for row in rows:
        if row[2] == "↓FoodID/→ParameterID":
            header = row
            break
    if header is None:
        raise SystemExit("Could not find the ParameterID header row in Data_Table")
    col = {pid: i for i, pid in enumerate(header) if isinstance(pid, int)}
    missing = [k for k, pid in PARAMS.items() if pid not in col]
    if missing:
        raise SystemExit(f"Parameters not found in Data_Table: {missing}")

    foods = []
    for row in rows:
        if not isinstance(row[2], (int, float)):
            continue
        food_id = int(row[2])
        foods.append(
            [food_id, row[0] or "", row[1] or "", groups.get(food_id, "")]
            + [_num(row[col[pid]]) for pid in PARAMS.values()]
        )

    attribution = (
        f"© The Danish Food Composition Database (http://fcdb.fooddata.dk), "
        f"version {version}, {date}. DTU Food Institute"
    )
    return {
        "source": "frida",
        "version": version,
        "date": date,
        "license": "CC BY 4.0",
        "license_url": "https://creativecommons.org/licenses/by/4.0/",
        "attribution": attribution,
        "url": "https://frida.fooddata.dk",
        "fields": FIELDS,
        "foods": foods,
    }


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    data = convert(Path(sys.argv[1]), sys.argv[2], sys.argv[3])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {len(data['foods'])} foods to {OUT}")


if __name__ == "__main__":
    main()
